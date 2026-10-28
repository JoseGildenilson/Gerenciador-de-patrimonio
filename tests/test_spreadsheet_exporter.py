"""Testes unitários para o gerador de planilhas Excel de ativos patrimoniais."""

import io
from datetime import datetime, timezone, date
from decimal import Decimal
import openpyxl
from src.domain.models import Asset, Category, Location, AssetMaintenance
from src.infrastructure.spreadsheet_exporter import (
    HEADERS,
    COLUMN_WIDTHS,
    _format_maintenances,
    _sanitize_formula,
    generate_assets_spreadsheet,
)


def test_sanitize_formula():
    assert _sanitize_formula("=1+1") == "'=1+1"
    assert _sanitize_formula("+cmd") == "'+cmd"
    assert _sanitize_formula("-10") == "'-10"
    assert _sanitize_formula("@SUM") == "'@SUM"
    assert _sanitize_formula("Notebook") == "Notebook"
    assert _sanitize_formula(None) is None


def test_format_maintenances_empty():
    asset = Asset(id=1, description="Notebook Test")
    asset.maintenances = []
    assert _format_maintenances(asset) == ""


def test_format_maintenances_with_data():
    asset = Asset(id=1, description="Notebook Test")
    m1 = AssetMaintenance(
        id=1,
        asset_id=1,
        maintenance_type="preventiva",
        description="Troca de pasta térmica",
        cost=Decimal("150.00"),
        maintenance_date=datetime(2026, 5, 10, tzinfo=timezone.utc),
    )
    m2 = AssetMaintenance(
        id=2,
        asset_id=1,
        maintenance_type="corretiva",
        description="Troca de tela",
        cost=None,
        maintenance_date=datetime(2026, 6, 15, tzinfo=timezone.utc),
    )
    asset.maintenances = [m2, m1]

    formatted = _format_maintenances(asset)
    assert "PREVENTIVA: Troca de pasta térmica (R$ 150.00)" in formatted
    assert "CORRETIVA: Troca de tela (R$ 0,00)" in formatted


def test_generate_assets_spreadsheet():
    cat = Category(id=1, name="Portáteis")
    loc = Location(id=1, building="Prédio 1", floor="Térreo", sector="TI")
    asset = Asset(
        id=1,
        patrimony_number="PAT-EXP-001",
        serial_number="SN-EXP-001",
        description="Dell Inspiron 15",
        brand="Dell",
        model="Inspiron 15",
        purchase_cost=Decimal("4500.00"),
        purchase_date=date(2026, 1, 10),
        status="em_uso",
        specs={"ram": "16GB"},
        category=cat,
        location=loc,
        maintenances=[],
    )

    data = generate_assets_spreadsheet([asset])
    assert isinstance(data, bytes)
    assert len(data) > 0

    # Carrega planilha gerada
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.title == "Ativos Patrimoniais"
    assert ws.cell(row=1, column=1).value == "CATEGORIA"
    assert ws.cell(row=2, column=1).value == "Portáteis"
    assert ws.cell(row=2, column=5).value == "PAT-EXP-001"
