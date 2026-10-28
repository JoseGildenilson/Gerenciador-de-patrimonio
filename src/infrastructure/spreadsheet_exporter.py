"""Utilitário para geração e exportação de planilhas Excel (.xlsx) de ativos patrimoniais."""

import io
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from src.domain.models import Asset


HEADERS = [
    "CATEGORIA",
    "LOCALIZAÇÃO",
    "STATUS",
    "DESCRIÇÃO",
    "Nº PATRIMÔNIO",
    "Nº SÉRIE",
    "MARCA",
    "MODELO",
    "CUSTO AQUISIÇÃO",
    "DATA AQUISIÇÃO",
    "HISTÓRICO DE MANUTENÇÕES",
    "ESPECIFICAÇÕES",
]

COLUMN_WIDTHS = {
    1: 22.0,  # A - CATEGORIA
    2: 24.0,  # B - LOCALIZAÇÃO
    3: 16.0,  # C - STATUS
    4: 38.0,  # D - DESCRIÇÃO
    5: 18.0,  # E - Nº PATRIMÔNIO
    6: 18.0,  # F - Nº SÉRIE
    7: 16.0,  # G - MARCA
    8: 18.0,  # H - MODELO
    9: 18.0,  # I - CUSTO AQUISIÇÃO
    10: 16.0,  # J - DATA AQUISIÇÃO
    11: 45.0,  # K - HISTÓRICO DE MANUTENÇÕES
    12: 40.0,  # L - ESPECIFICAÇÕES
}


def _format_maintenances(asset: Asset) -> str:
    """Formata as manutenções cadastradas em ordem cronológica crescente com quebras de linha."""
    if not asset.maintenances:
        return ""

    sorted_maints = sorted(
        asset.maintenances,
        key=lambda m: m.maintenance_date or datetime.min.replace(tzinfo=timezone.utc),
    )

    lines: list[str] = []
    for m in sorted_maints:
        date_str = m.maintenance_date.strftime("%d/%m/%Y") if m.maintenance_date else ""
        type_str = m.maintenance_type or "Geral"
        desc_str = m.description or "Sem descrição"
        cost_str = f"R$ {m.cost:.2f}" if m.cost is not None else "R$ 0,00"
        lines.append(f"[{date_str}] {type_str.upper()}: {desc_str} ({cost_str})")

    return "\n".join(lines)


def _sanitize_formula(val: str | None) -> str | None:
    """Higieniza valores de texto para evitar ataques de injeção de fórmulas CSV/Excel (CWE-1236)."""
    if val is None:
        return None
    val_str = str(val).strip()
    if val_str.startswith(("=", "+", "-", "@")):
        return "'" + val_str
    return val_str


def generate_assets_spreadsheet(assets: list[Asset]) -> bytes:
    """Gera um arquivo de planilha Excel (.xlsx) contendo a relação de ativos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ativos Patrimoniais"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Calibri", size=10)
    data_alignment = Alignment(vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.row_dimensions[1].height = 28.0
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(col_idx, 15.0)

    for row_idx, asset in enumerate(assets, start=2):
        loc_str = f"{asset.location.building} - {asset.location.sector}" if asset.location else ""
        cost_str = f"R$ {asset.purchase_cost:.2f}" if asset.purchase_cost is not None else ""
        date_str = asset.purchase_date.strftime("%d/%m/%Y") if asset.purchase_date else ""
        specs_str = str(asset.specs) if asset.specs else ""

        row_data = [
            _sanitize_formula(asset.category.name if asset.category else None),
            _sanitize_formula(loc_str),
            _sanitize_formula(asset.status),
            _sanitize_formula(asset.description),
            _sanitize_formula(asset.patrimony_number),
            _sanitize_formula(asset.serial_number),
            _sanitize_formula(asset.brand),
            _sanitize_formula(asset.model),
            cost_str,
            date_str,
            _sanitize_formula(_format_maintenances(asset)),
            _sanitize_formula(specs_str),
        ]

        ws.row_dimensions[row_idx].height = 20.0
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = thin_border

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
